import datetime
import json
import logging
import os
import re
from typing import Any

import database

_log = logging.getLogger("broker_database")


def clean_field(val: Any) -> str:
    """Sanitizes string inputs, removing extra whitespace and trailing float decimals."""
    if val is None:
        return ""
    s = str(val).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return re.sub(r"\s+", " ", s)


def normalize_entity_key(name: Any, reg_no: Any = "") -> str:
    """
    Generates a canonical deduplication key for a broker entity.
    If a valid SEBI registration number is present, uses REG_<REG_NO>.
    Otherwise normalizes legal company entity suffixes.
    """
    clean_reg = clean_field(reg_no)
    if clean_reg:
        r = re.sub(r"[^A-Z0-9]", "", clean_reg.upper())
        if r:
            return f"REG_{r}"

    s = clean_field(name).upper()
    if not s:
        return ""

    # Strip common corporate designations for name-based matching
    suffixes = [
        "PRIVATE LIMITED",
        "PVT LTD",
        "PVT. LTD.",
        "PVT.LTD.",
        "LIMITED",
        "LTD",
        "LLP",
        "INCORPORATED",
        "INC",
        "CORP",
    ]
    for suf in suffixes:
        s = s.replace(suf, "")

    norm = re.sub(r"[^A-Z0-9]", "", s)
    return norm


def parse_xls_brokers(file_path: str) -> list[dict[str, Any]]:
    """Parses SEBI registered stock brokers XLS spreadsheet files."""
    records: list[dict[str, Any]] = []
    if not os.path.exists(file_path):
        return records

    try:
        import xlrd

        wb = xlrd.open_workbook(file_path)
        sheet = wb.sheet_by_index(0)
        fname = os.path.basename(file_path)

        seg_match = re.search(r"Brokers in\s+(.*?)\s+Segment", fname, re.IGNORECASE)
        segment = seg_match.group(1).strip() if seg_match else "General"

        # Headers are on row index 2, data starts at row index 3
        for r in range(3, sheet.nrows):
            row = [sheet.cell_value(r, c) for c in range(sheet.ncols)]
            name = clean_field(row[0]) if len(row) > 0 else ""
            reg_no = clean_field(row[1]) if len(row) > 1 else ""
            contact = clean_field(row[2]) if len(row) > 2 else ""
            addr = clean_field(row[3]) if len(row) > 3 else ""
            email = clean_field(row[4]) if len(row) > 4 else ""
            phone = clean_field(row[5]) if len(row) > 5 else ""
            fax = clean_field(row[6]) if len(row) > 6 else ""
            city = clean_field(row[7]) if len(row) > 7 else ""
            state = clean_field(row[8]) if len(row) > 8 else ""
            pincode = clean_field(row[9]) if len(row) > 9 else ""
            v_from = clean_field(row[17]) if len(row) > 17 else ""
            v_to = clean_field(row[18]) if len(row) > 18 else ""
            exch = clean_field(row[20]) if len(row) > 20 else ""
            tname = clean_field(row[21]) if len(row) > 21 else ""

            if name or reg_no or tname:
                records.append(
                    {
                        "name": name or tname,
                        "registration_no": reg_no,
                        "trade_name": tname,
                        "contact_person": contact,
                        "email": email,
                        "phone": phone,
                        "fax": fax,
                        "address": addr,
                        "city": city,
                        "state": state,
                        "pincode": pincode,
                        "exchange": exch,
                        "valid_from": v_from,
                        "valid_to": v_to,
                        "segments": [segment] if segment else [],
                        "sources": [fname],
                    }
                )
    except Exception as e:
        _log.error("Error parsing XLS file %s: %s", file_path, e)

    return records


def parse_xlsx_brokers(file_path: str) -> list[dict[str, Any]]:
    """Parses indian_brokers.xlsx spreadsheet file."""
    records: list[dict[str, Any]] = []
    if not os.path.exists(file_path):
        return records

    try:
        import openpyxl

        wb = openpyxl.load_workbook(file_path, read_only=True)
        fname = os.path.basename(file_path)

        for sheetname in wb.sheetnames:
            sheet = wb[sheetname]
            for row in sheet.iter_rows(values_only=True):
                if row and row[0]:
                    nm = clean_field(row[0])
                    if nm and nm.upper() not in (
                        "TRADE NAME",
                        "MEMBER NAME",
                        "INDIAN STOCK MARKET REGISTERED BROKERS",
                    ):
                        records.append(
                            {
                                "name": nm,
                                "segments": [sheetname],
                                "sources": [fname],
                            }
                        )
    except Exception as e:
        _log.error("Error parsing XLSX file %s: %s", file_path, e)

    return records


def parse_txt_brokers(file_path: str) -> list[dict[str, Any]]:
    """Parses text list broker files."""
    records: list[dict[str, Any]] = []
    if not os.path.exists(file_path):
        return records

    try:
        fname = os.path.basename(file_path)
        with open(file_path, encoding="utf-8", errors="ignore") as f:
            for line in f:
                nm = clean_field(line)
                if nm and nm.upper() not in (
                    "TRADE NAME",
                    "MEMBER NAME",
                    "INDIAN STOCK MARKET REGISTERED BROKERS",
                ):
                    records.append(
                        {
                            "name": nm,
                            "sources": [fname],
                        }
                    )
    except Exception as e:
        _log.error("Error parsing TXT file %s: %s", file_path, e)

    return records


def parse_pdf_brokers(file_path: str) -> list[dict[str, Any]]:
    """Parses NSE Clearing / NCL registered members PDF files."""
    records: list[dict[str, Any]] = []
    if not os.path.exists(file_path):
        return records

    try:
        import pdfplumber

        fname = os.path.basename(file_path)
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                tables = page.extract_tables()
                for t in tables:
                    for r in t:
                        if not r or len(r) < 3:
                            continue
                        sr, name, code = r[0], r[1], r[2]
                        seg = r[3] if len(r) > 3 else ""
                        if sr and ("Sr" in str(sr) or "No" in str(sr)):
                            continue
                        clean_nm = clean_field(name)
                        clean_cd = clean_field(code)
                        clean_sg = clean_field(seg)
                        if clean_nm:
                            records.append(
                                {
                                    "name": clean_nm,
                                    "member_codes": [clean_cd] if clean_cd else [],
                                    "segments": [clean_sg] if clean_sg else [],
                                    "sources": [fname],
                                }
                            )
    except Exception as e:
        _log.error("Error parsing PDF file %s: %s", file_path, e)

    return records


def ingest_and_deduplicate_all_brokers(
    brokers_dir: str = "Brokers",
) -> dict[str, dict[str, Any]]:
    """
    Ingests all 11 broker datasets in the target folder, deduplicates entries,
    and cross-verifies records across sources using dual name and SEBI registration indexing.
    Returns dict mapping norm_key -> aggregated broker record.
    """
    all_brokers: dict[str, dict[str, Any]] = {}
    name_to_key: dict[str, str] = {}
    reg_to_key: dict[str, str] = {}

    def merge_record(rec: dict[str, Any]) -> None:
        nm = rec.get("name", "")
        reg = rec.get("registration_no", "")
        tname = rec.get("trade_name", "")
        primary = nm or tname
        if not primary and not reg:
            return

        name_key = normalize_entity_key(primary, reg_no="")
        reg_key = re.sub(r"[^A-Z0-9]", "", clean_field(reg).upper()) if reg else ""

        # Determine target canonical key
        target_key = None
        if reg_key and reg_key in reg_to_key:
            target_key = reg_to_key[reg_key]
        elif name_key and name_key in name_to_key:
            target_key = name_to_key[name_key]
        else:
            target_key = f"REG_{reg_key}" if reg_key else name_key

        if not target_key:
            return

        # If name_key and reg_key exist under different canonical keys, merge records
        if (
            name_key
            and reg_key
            and name_key in name_to_key
            and reg_to_key.get(reg_key)
            and name_to_key[name_key] != reg_to_key[reg_key]
        ):
            old_key = name_to_key[name_key]
            new_key = reg_to_key[reg_key]
            if old_key in all_brokers and new_key in all_brokers:
                old_b = all_brokers.pop(old_key)
                new_b = all_brokers[new_key]
                if not new_b["name"] and old_b["name"]:
                    new_b["name"] = old_b["name"]
                if not new_b["registration_no"] and old_b["registration_no"]:
                    new_b["registration_no"] = old_b["registration_no"]
                if not new_b["trade_name"] and old_b["trade_name"]:
                    new_b["trade_name"] = old_b["trade_name"]
                if not new_b["contact_person"] and old_b["contact_person"]:
                    new_b["contact_person"] = old_b["contact_person"]
                if not new_b["email"] and old_b["email"]:
                    new_b["email"] = old_b["email"]
                if not new_b["phone"] and old_b["phone"]:
                    new_b["phone"] = old_b["phone"]
                if not new_b["fax"] and old_b["fax"]:
                    new_b["fax"] = old_b["fax"]
                if not new_b["address"] and old_b["address"]:
                    new_b["address"] = old_b["address"]
                if not new_b["city"] and old_b["city"]:
                    new_b["city"] = old_b["city"]
                if not new_b["state"] and old_b["state"]:
                    new_b["state"] = old_b["state"]
                if not new_b["pincode"] and old_b["pincode"]:
                    new_b["pincode"] = old_b["pincode"]
                if not new_b["exchange"] and old_b["exchange"]:
                    new_b["exchange"] = old_b["exchange"]
                if not new_b["valid_from"] and old_b["valid_from"]:
                    new_b["valid_from"] = old_b["valid_from"]
                if not new_b["valid_to"] and old_b["valid_to"]:
                    new_b["valid_to"] = old_b["valid_to"]
                new_b["member_codes"].update(old_b["member_codes"])
                new_b["segments"].update(old_b["segments"])
                new_b["sources"].update(old_b["sources"])

                # Redirect index pointers
                for k, v in list(name_to_key.items()):
                    if v == old_key:
                        name_to_key[k] = new_key
                for k, v in list(reg_to_key.items()):
                    if v == old_key:
                        reg_to_key[k] = new_key
                target_key = new_key

        if target_key not in all_brokers:
            all_brokers[target_key] = {
                "norm_key": target_key,
                "name": clean_field(primary),
                "registration_no": clean_field(reg),
                "trade_name": clean_field(tname),
                "contact_person": clean_field(rec.get("contact_person", "")),
                "email": clean_field(rec.get("email", "")),
                "phone": clean_field(rec.get("phone", "")),
                "fax": clean_field(rec.get("fax", "")),
                "address": clean_field(rec.get("address", "")),
                "city": clean_field(rec.get("city", "")),
                "state": clean_field(rec.get("state", "")),
                "pincode": clean_field(rec.get("pincode", "")),
                "exchange": clean_field(rec.get("exchange", "")),
                "valid_from": clean_field(rec.get("valid_from", "")),
                "valid_to": clean_field(rec.get("valid_to", "")),
                "member_codes": set(rec.get("member_codes", [])),
                "segments": set(rec.get("segments", [])),
                "sources": set(rec.get("sources", [])),
            }
        else:
            b = all_brokers[target_key]
            if not b["name"] and primary:
                b["name"] = clean_field(primary)
            if not b["registration_no"] and reg:
                b["registration_no"] = clean_field(reg)
            if not b["trade_name"] and tname:
                b["trade_name"] = clean_field(tname)
            if not b["contact_person"] and rec.get("contact_person"):
                b["contact_person"] = clean_field(rec.get("contact_person", ""))
            if not b["email"] and rec.get("email"):
                b["email"] = clean_field(rec.get("email", ""))
            if not b["phone"] and rec.get("phone"):
                b["phone"] = clean_field(rec.get("phone", ""))
            if not b["fax"] and rec.get("fax"):
                b["fax"] = clean_field(rec.get("fax", ""))
            if not b["address"] and rec.get("address"):
                b["address"] = clean_field(rec.get("address", ""))
            if not b["city"] and rec.get("city"):
                b["city"] = clean_field(rec.get("city", ""))
            if not b["state"] and rec.get("state"):
                b["state"] = clean_field(rec.get("state", ""))
            if not b["pincode"] and rec.get("pincode"):
                b["pincode"] = clean_field(rec.get("pincode", ""))
            if not b["exchange"] and rec.get("exchange"):
                b["exchange"] = clean_field(rec.get("exchange", ""))
            if not b["valid_from"] and rec.get("valid_from"):
                b["valid_from"] = clean_field(rec.get("valid_from", ""))
            if not b["valid_to"] and rec.get("valid_to"):
                b["valid_to"] = clean_field(rec.get("valid_to", ""))

            if rec.get("member_codes"):
                b["member_codes"].update(rec.get("member_codes", []))
            if rec.get("segments"):
                b["segments"].update(rec.get("segments", []))
            if rec.get("sources"):
                b["sources"].update(rec.get("sources", []))

        if name_key:
            name_to_key[name_key] = target_key
        if reg_key:
            reg_to_key[reg_key] = target_key

    if not os.path.exists(brokers_dir):
        _log.warning("Brokers directory %s does not exist", brokers_dir)
        return all_brokers

    for root, _, files in os.walk(brokers_dir):
        for f in sorted(files):
            fpath = os.path.join(root, f)
            if f.endswith(".xls"):
                recs = parse_xls_brokers(fpath)
            elif f.endswith(".xlsx"):
                recs = parse_xlsx_brokers(fpath)
            elif f.endswith(".txt"):
                recs = parse_txt_brokers(fpath)
            elif f.endswith(".pdf"):
                recs = parse_pdf_brokers(fpath)
            else:
                continue

            for r in recs:
                merge_record(r)

    # Compute verification status
    for b in all_brokers.values():
        num_sources = len(b["sources"])
        has_reg = bool(b["registration_no"])
        has_mcode = bool(b["member_codes"])

        if has_reg or (num_sources >= 2 and has_mcode):
            b["verification_status"] = "SEBI_VERIFIED"
        elif num_sources >= 2 or has_mcode:
            b["verification_status"] = "MULTI_SOURCE_VERIFIED"
        else:
            b["verification_status"] = "SINGLE_SOURCE"

    return all_brokers


def build_and_seed_broker_database(
    brokers_dir: str = "Brokers", force_rebuild: bool = False
) -> int:
    """
    Parses all broker files, deduplicates entries, cross-verifies, and populates
    the `registered_brokers` SQLite table.
    """
    database.init_db()

    conn = database.get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT COUNT(*) FROM registered_brokers")
    existing_count = cursor.fetchone()[0]
    conn.close()

    if existing_count > 0 and not force_rebuild:
        _log.info("Registered brokers table already populated (%d records)", existing_count)
        return int(existing_count)

    _log.info("Parsing and ingesting broker data from %s...", brokers_dir)
    all_brokers = ingest_and_deduplicate_all_brokers(brokers_dir)

    now_str = datetime.datetime.now().isoformat()
    rows_to_insert = []

    for b in all_brokers.values():
        rows_to_insert.append(
            (
                b["norm_key"],
                b["name"],
                b["registration_no"],
                b["trade_name"],
                b["contact_person"],
                b["email"],
                b["phone"],
                b["fax"],
                b["address"],
                b["city"],
                b["state"],
                b["pincode"],
                b["exchange"],
                b["valid_from"],
                b["valid_to"],
                json.dumps(sorted(b["member_codes"])),
                json.dumps(sorted(b["segments"])),
                json.dumps(sorted(b["sources"])),
                b["verification_status"],
                now_str,
            )
        )

    if force_rebuild:
        database._execute_with_retry("DELETE FROM registered_brokers")

    query = """
    INSERT INTO registered_brokers (
        norm_key, name, registration_no, trade_name, contact_person,
        email, phone, fax, address, city, state, pincode, exchange,
        valid_from, valid_to, member_codes, segments, sources,
        verification_status, updated_at
    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ON CONFLICT(norm_key) DO UPDATE SET
        name=excluded.name,
        registration_no=CASE WHEN excluded.registration_no != '' THEN excluded.registration_no ELSE registered_brokers.registration_no END,
        trade_name=CASE WHEN excluded.trade_name != '' THEN excluded.trade_name ELSE registered_brokers.trade_name END,
        contact_person=CASE WHEN excluded.contact_person != '' THEN excluded.contact_person ELSE registered_brokers.contact_person END,
        email=CASE WHEN excluded.email != '' THEN excluded.email ELSE registered_brokers.email END,
        phone=CASE WHEN excluded.phone != '' THEN excluded.phone ELSE registered_brokers.phone END,
        fax=CASE WHEN excluded.fax != '' THEN excluded.fax ELSE registered_brokers.fax END,
        address=CASE WHEN excluded.address != '' THEN excluded.address ELSE registered_brokers.address END,
        city=CASE WHEN excluded.city != '' THEN excluded.city ELSE registered_brokers.city END,
        state=CASE WHEN excluded.state != '' THEN excluded.state ELSE registered_brokers.state END,
        pincode=CASE WHEN excluded.pincode != '' THEN excluded.pincode ELSE registered_brokers.pincode END,
        exchange=CASE WHEN excluded.exchange != '' THEN excluded.exchange ELSE registered_brokers.exchange END,
        valid_from=CASE WHEN excluded.valid_from != '' THEN excluded.valid_from ELSE registered_brokers.valid_from END,
        valid_to=CASE WHEN excluded.valid_to != '' THEN excluded.valid_to ELSE registered_brokers.valid_to END,
        member_codes=excluded.member_codes,
        segments=excluded.segments,
        sources=excluded.sources,
        verification_status=excluded.verification_status,
        updated_at=excluded.updated_at
    """

    with database.get_connection() as conn:
        cursor = conn.cursor()
        cursor.executemany(query, rows_to_insert)
        conn.commit()

    conn = database.get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT COUNT(*) FROM registered_brokers")
    final_count = cursor.fetchone()[0]
    conn.close()

    _log.info("Successfully seeded registered_brokers database with %d records", final_count)
    return int(final_count)


def get_registered_broker(query_str: str) -> dict[str, Any] | None:
    """Retrieves a single registered broker by name, registration number, or norm_key."""
    database.init_db()
    cleaned = clean_field(query_str)
    if not cleaned:
        return None

    norm = normalize_entity_key(cleaned)
    norm_reg = normalize_entity_key("", cleaned)

    query = """
    SELECT * FROM registered_brokers
    WHERE norm_key = ? OR norm_key = ? OR LOWER(name) = LOWER(?) OR LOWER(registration_no) = LOWER(?)
    LIMIT 1
    """
    row = database._fetch_with_retry(
        query, (norm, norm_reg, cleaned, cleaned), fetch_all=False
    )
    if row:
        res = dict(row)
        res["member_codes"] = json.loads(res.get("member_codes", "[]"))
        res["segments"] = json.loads(res.get("segments", "[]"))
        res["sources"] = json.loads(res.get("sources", "[]"))
        return res
    return None


def search_registered_brokers(
    keyword: str = "", segment: str = "", limit: int = 50
) -> list[dict[str, Any]]:
    """Searches registered brokers by keyword or segment filter."""
    database.init_db()
    kw = f"%{clean_field(keyword)}%"
    seg = f"%{clean_field(segment)}%"

    if keyword and segment:
        query = """
        SELECT * FROM registered_brokers
        WHERE (name LIKE ? OR registration_no LIKE ? OR trade_name LIKE ?)
          AND segments LIKE ?
        ORDER BY name ASC LIMIT ?
        """
        rows = database._fetch_with_retry(query, (kw, kw, kw, seg, limit), fetch_all=True)
    elif keyword:
        query = """
        SELECT * FROM registered_brokers
        WHERE name LIKE ? OR registration_no LIKE ? OR trade_name LIKE ?
        ORDER BY name ASC LIMIT ?
        """
        rows = database._fetch_with_retry(query, (kw, kw, kw, limit), fetch_all=True)
    elif segment:
        query = """
        SELECT * FROM registered_brokers
        WHERE segments LIKE ?
        ORDER BY name ASC LIMIT ?
        """
        rows = database._fetch_with_retry(query, (seg, limit), fetch_all=True)
    else:
        query = "SELECT * FROM registered_brokers ORDER BY name ASC LIMIT ?"
        rows = database._fetch_with_retry(query, (limit,), fetch_all=True)

    results = []
    if rows:
        for r in rows:
            res = dict(r)
            res["member_codes"] = json.loads(res.get("member_codes", "[]"))
            res["segments"] = json.loads(res.get("segments", "[]"))
            res["sources"] = json.loads(res.get("sources", "[]"))
            results.append(res)
    return results


def get_registered_broker_stats() -> dict[str, Any]:
    """Retrieves high-level analytical statistics from registered_brokers table."""
    database.init_db()
    with database.get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT COUNT(*) FROM registered_brokers")
        total = cursor.fetchone()[0]

        cursor.execute(
            "SELECT verification_status, COUNT(*) FROM registered_brokers GROUP BY verification_status"
        )
        status_counts = dict(cursor.fetchall())

    return {
        "total_unique_brokers": total,
        "verification_breakdown": status_counts,
    }


if __name__ == "__main__":
    count = build_and_seed_broker_database("Brokers", force_rebuild=True)
    stats = get_registered_broker_stats()
    print(f"Ingestion & Database Seeding Complete: {count} brokers stored.")
    print("Statistics:", stats)
